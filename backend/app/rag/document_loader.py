"""Document loading utilities for the ABES RAG knowledge base."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Callable

from langchain_core.documents import Document

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".md", ".csv", ".json",
    ".xlsx", ".pptx", ".html", ".htm",
}


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _document(path: Path, text: str, file_type: str, **metadata: object) -> Document:
    return Document(
        page_content=text,
        metadata={"source": str(path), "file_type": file_type, **metadata},
    )


def load_text(file_path: Path) -> list[Document]:
    return [_document(file_path, _read_text(file_path), "text")]


def load_markdown(file_path: Path) -> list[Document]:
    return [_document(file_path, _read_text(file_path), "markdown")]


def load_csv(file_path: Path) -> list[Document]:
    documents: list[Document] = []
    with file_path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        for row_number, row in enumerate(reader, start=1):
            fields = [
                f"{key}: {str(value).strip()}"
                for key, value in row.items()
                if value is not None and str(value).strip()
            ]
            if fields:
                documents.append(
                    _document(file_path, "\n".join(fields), "csv", row=row_number)
                )
    return documents


def load_json(file_path: Path) -> list[Document]:
    data = json.loads(_read_text(file_path))
    return [_document(file_path, json.dumps(data, indent=2, ensure_ascii=False), "json")]


def load_pdf(file_path: Path) -> list[Document]:
    from pypdf import PdfReader

    documents: list[Document] = []
    reader = PdfReader(str(file_path))
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            documents.append(_document(file_path, text, "pdf", page=page_number))
    return documents


def load_docx(file_path: Path) -> list[Document]:
    from docx import Document as DocxDocument

    document = DocxDocument(str(file_path))
    sections: list[str] = []
    sections.extend(p.text.strip() for p in document.paragraphs if p.text.strip())
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                sections.append(" | ".join(values))
    content = "\n".join(sections)
    return [_document(file_path, content, "docx")] if content.strip() else []


def load_xlsx(file_path: Path) -> list[Document]:
    from openpyxl import load_workbook

    documents: list[Document] = []
    workbook = load_workbook(str(file_path), read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            rows = list(worksheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [
                str(value).strip() if value is not None else f"column_{i + 1}"
                for i, value in enumerate(rows[0])
            ]
            for row_number, row in enumerate(rows[1:], start=2):
                fields = []
                for index, value in enumerate(row):
                    if value is None or not str(value).strip():
                        continue
                    header = headers[index] if index < len(headers) else f"column_{index + 1}"
                    fields.append(f"{header}: {str(value).strip()}")
                if fields:
                    documents.append(
                        _document(
                            file_path,
                            "\n".join(fields),
                            "xlsx",
                            sheet=worksheet.title,
                            row=row_number,
                        )
                    )
    finally:
        workbook.close()
    return documents


def load_pptx(file_path: Path) -> list[Document]:
    from pptx import Presentation

    documents: list[Document] = []
    presentation = Presentation(str(file_path))
    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        content = "\n".join(texts)
        if content:
            documents.append(_document(file_path, content, "pptx", slide=slide_number))
    return documents


def load_html(file_path: Path) -> list[Document]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_read_text(file_path), "html.parser")
    for element in soup(["script", "style", "noscript"]):
        element.decompose()
    text = soup.get_text(separator="\n", strip=True)
    return [_document(file_path, text, "html")] if text else []


_LOADERS: dict[str, Callable[[Path], list[Document]]] = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".txt": load_text,
    ".md": load_markdown,
    ".csv": load_csv,
    ".json": load_json,
    ".xlsx": load_xlsx,
    ".pptx": load_pptx,
    ".html": load_html,
    ".htm": load_html,
}


def load_file(file_path: Path) -> list[Document]:
    """Load a supported file into LangChain documents."""
    if not file_path.is_file():
        raise FileNotFoundError(f"File not found: {file_path}")
    loader = _LOADERS.get(file_path.suffix.lower())
    if loader is None:
        raise ValueError(f"Unsupported file type: {file_path.suffix}")
    return loader(file_path)
